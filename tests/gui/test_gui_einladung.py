"""TST-0028 – Seite „Einladung“ und Word-Export in der GUI (CON-0019)."""

from datetime import date, time

import docx
from PySide6.QtCore import QDate

from konfiplaner.anwendung import Sitzung
from konfiplaner.einstellungen import YamlEinstellungen
from konfiplaner.gui.hauptfenster import SEITE_EINLADUNG, Hauptfenster
from konfiplaner.tagesplan import Block, Tagesplan
from tests.gui.test_gui_ablaeufe import FakeDialoge
from tests.unit.test_auswahl_in_der_sitzung import MI, ZEILEN
from tests.xlsx_helfer import lehrer_mit_terminen, schreibe


def fenster(qtbot, tmp_path, laden=True):
    d = FakeDialoge()
    d.datei = schreibe(
        tmp_path / "k.xlsx",
        lehrer_mit_terminen(ZEILEN, ("Deutsch", *MI), ("Sport", *MI), ("Kunst", "Mi", "15:00", "16:00")),
    )
    f = Hauptfenster(Sitzung(YamlEinstellungen(tmp_path / "e.yaml")), d, auto_laden=False)
    qtbot.addWidget(f)
    f.show()
    if laden:
        f.datei_oeffnen()
        f.reiter.setCurrentIndex(SEITE_EINLADUNG)
    return f, d


def test_schritt_ist_ohne_datei_gesperrt_und_danach_frei(qtbot, tmp_path):
    f, _ = fenster(qtbot, tmp_path, laden=False)
    assert not f.reiter.isTabEnabled(SEITE_EINLADUNG)
    f.datei_oeffnen()
    assert f.reiter.isTabEnabled(SEITE_EINLADUNG)
    assert f.leiste.schritte[SEITE_EINLADUNG].titel == "Einladung"


def test_standardblocks_stehen_in_der_liste(qtbot, tmp_path):
    f, _ = fenster(qtbot, tmp_path)
    s = f.einladung
    assert s.liste.count() == 9 and s.liste.item(0).text().startswith("ab 8.00 Uhr")
    assert s.liste.currentRow() == 0 and s.block_titel.text() == "Stehkaffee im Foyer"


def test_kopf_wird_gespeichert(qtbot, tmp_path):
    f, _ = fenster(qtbot, tmp_path)
    s = f.einladung
    s.titel.setText("Mein Tag")
    s.thema.setText("Methoden")
    s.ziele.setPlainText("Eins\n\nZwei")
    s.datum.setDate(QDate(2026, 11, 9))
    plan = f.sitzung.tagesplan
    assert (plan.titel, plan.thema, plan.ziele, plan.datum) == (
        "Mein Tag",
        "Methoden",
        ("Eins", "Zwei"),
        date(2026, 11, 9),
    )
    neu = Sitzung(YamlEinstellungen(tmp_path / "e.yaml"))
    assert neu.tagesplan.titel == "Mein Tag"


def test_bloecke_hinzufuegen_loeschen_verschieben(qtbot, tmp_path):
    f, _ = fenster(qtbot, tmp_path)
    s = f.einladung
    s.liste.setCurrentRow(1)
    s.neu.click()
    assert s.liste.count() == 10 and s.liste.currentRow() == 2 and s.block_titel.text() == "Neuer Block"
    s.hoch.click()
    assert s.liste.currentRow() == 1 and f.sitzung.tagesplan.bloecke[1].titel == "Neuer Block"
    s.runter.click()
    s.loeschen.click()
    assert s.liste.count() == 9 and all(b.titel != "Neuer Block" for b in f.sitzung.tagesplan.bloecke)


def test_blockfelder_aendern_den_plan(qtbot, tmp_path):
    f, _ = fenster(qtbot, tmp_path)
    s = f.einladung
    s.liste.setCurrentRow(1)
    s.block_titel.setText("Willkommen")
    s.block_wer.setText("Leitung")
    s.konferenzblock.setChecked(True)
    b = f.sitzung.tagesplan.bloecke[1]
    assert (b.titel, b.wer, b.art) == ("Willkommen", "Leitung", "konferenzen")
    assert s.liste.item(1).text() == "8.30 – 8.40 Uhr  Willkommen"


def test_fehlerhafter_block_sperrt_den_export_ueberlappung_nicht(qtbot, tmp_path):
    f, _ = fenster(qtbot, tmp_path)
    s = f.einladung
    f.sitzung.setze_tagesplan(
        Tagesplan(
            bloecke=(
                Block(time(9), time(10), "bereich", "A"),
                Block(time(9, 30), time(10, 30), "bereich", "B"),
            )
        )
    )
    s.aktualisieren()
    assert "überschneidet" in s.pruefung.text() and s.export.isEnabled()
    f.sitzung.setze_tagesplan(Tagesplan(bloecke=(Block(time(10), time(9), "bereich", "A"),)))
    s.aktualisieren()
    assert "Ende" in s.pruefung.text() and not s.export.isEnabled()


def test_export_schreibt_die_gewaehlten_konferenzen(qtbot, tmp_path):
    f, d = fenster(qtbot, tmp_path)
    d.ziel = tmp_path / "Einladung.docx"
    f.einladung.exportieren()
    dok = docx.Document(str(d.ziel))
    zeile = next(r for r in dok.tables[0].rows if "Fachkonferenzen" in r.cells[1].text)
    for name in f.sitzung.ausgewaehlte_konferenzen():
        assert name in zeile.cells[1].text
    assert "Einladung exportiert" in f.statusBar().currentMessage()


def test_ueberschreiben_nur_nach_bestaetigung(qtbot, tmp_path):
    f, d = fenster(qtbot, tmp_path)
    d.ziel = tmp_path / "E.docx"
    d.ziel.write_text("alt")
    d.antwort = False
    f.einladung.exportieren()
    assert d.ziel.read_text() == "alt" and "abgebrochen" in f.statusBar().currentMessage()
    d.antwort = True
    f.einladung.exportieren()
    assert docx.Document(str(d.ziel)).tables


def test_logo_waehlen_entfernen_und_fehlendes_logo(qtbot, tmp_path):
    f, d = fenster(qtbot, tmp_path)
    s = f.einladung
    assert s.logo_text.text() == "Kein Logo" and not s.logo_entfernen.isEnabled()
    d.logo = tmp_path / "fehlt.png"
    s.logo_waehlen.click()
    assert f.sitzung.logo == d.logo and "nicht gefunden" in s.logo_text.text()
    d.ziel = tmp_path / "E.docx"
    s.exportieren()
    assert "ohne Logo" in f.statusBar().currentMessage() and d.ziel.exists()
    s.logo_entfernen.click()
    assert f.sitzung.logo is None and s.logo_text.text() == "Kein Logo"
