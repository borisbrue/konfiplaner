"""TST-0016 – Konferenzplanung mit festen Zeitslots (CON-0013). Testnamen folgen den Szenarien der .feature-Datei."""

from datetime import time

from openpyxl import load_workbook

from konfiplaner.anwendung import Sitzung
from konfiplaner.excel_import import importiere
from konfiplaner.normalisierung import Normalisierung
from tests.xlsx_helfer import LEHRER, PERSONEN, THEMEN, lehrer_mit_terminen, schreibe

MI = ("Mi", "14:00", "15:00")
MI2 = ("Mi", "15:00", "16:00")

# Drei Lehrkräfte; Anna hat Deutsch (I) und Sport (II)
ZEILEN = [
    ["Eins", "Anna", "", "Deutsch", "Sport", None, None],
    ["Zwei", "Ben", "", "Deutsch", None, None, "Deutsch"],
    ["Drei", "Cem", "", "Sport", None, None, "Sport"],
]


def lese(tmp_path, zeilen, *termine, norm=None, **kw):
    return importiere(
        schreibe(tmp_path / "k.xlsx", lehrer_mit_terminen(zeilen, *termine, **kw)), normalisierung=norm
    )


def sitzung(tmp_path, zeilen, *termine, **kw):
    s = Sitzung()
    s.lade(schreibe(tmp_path / "k.xlsx", lehrer_mit_terminen(zeilen, *termine, **kw)))
    s.setze_profil(
        "festtermine"
    )  # SPEC-0008: Rangfolge bei gleichzeitigen Terminen (SPEC-0012: Profil wählbar)
    return s


def teilnehmer(plan, thema):
    return [n for t in plan.termine if t.thema == thema for n in t.teilnehmer]


def test_konferenztermine_ersetzen_das_zeitraster(tmp_path):
    erg = lese(tmp_path, ZEILEN, ("Deutsch", *MI), ("Sport", *MI2))
    assert erg.festtermine and not erg.fehler
    assert [(z.start, z.ende) for z in erg.zeitslots] == [(time(14), time(15)), (time(15), time(16))]
    assert erg.modell is not None and all(t.festtermine for t in erg.themen)


def test_konferenz_ohne_termin_ist_ein_fehler(tmp_path):
    erg = lese(tmp_path, ZEILEN, ("Deutsch", *MI))
    assert [f.text for f in erg.fehler] == ["Konferenz Sport hat keinen Termin im Blatt Konferenzen"]
    assert erg.modell is None


def test_konferenzname_wird_normalisiert(tmp_path):
    zeilen = [
        ["Eins", "Anna", "", "Deutsch", "Mathematik", None, None],
        ["Zwei", "Ben", "", "Deutsch", None, None, "Deutsch"],
        ["Drei", "Cem", "", "Mathematik", None, None, "Mathematik"],
    ]
    erg = lese(tmp_path, zeilen, ("Deutsch", *MI), ("Mathe", *MI2))
    fehler = [f for f in erg.fehler if f.blatt == "Konferenzen" and f.zeile == 3]
    assert len(fehler) == 1 and "Mathematik" in fehler[0].text and fehler[0].spalte == "Konferenz"
    n = Normalisierung()
    n.bestaetige("Mathe", ["Mathematik"])
    assert not lese(tmp_path, zeilen, ("Deutsch", *MI), ("Mathe", *MI2), norm=n).fehler


def test_ungueltige_zeit(tmp_path):
    erg = lese(tmp_path, ZEILEN, ("Deutsch", "Mi", "15:00", "14:00"), ("Sport", *MI2))
    assert any(f.spalte == "Ende" and f.zeile == 2 for f in erg.fehler)


def test_mehrere_termine_derselben_konferenz_ueberschneiden_sich(tmp_path):
    erg = lese(tmp_path, ZEILEN, ("Deutsch", *MI), ("Deutsch", "Mi", "14:30", "15:30"), ("Sport", *MI2))
    assert [(f.zeile, f.spalte) for f in erg.fehler] == [(3, "Start")] and "Zeile 2" in erg.fehler[0].text


def test_widerspruechliche_konferenzwerte(tmp_path):
    kopf = ["Konferenz", "Tag", "Start", "Ende", "Mindestanzahl"]
    erg = lese(
        tmp_path,
        ZEILEN,
        ("Deutsch", *MI, 1),
        ("Deutsch", "Mi", "16:00", "17:00", 2),
        ("Sport", *MI2, 1),
        kopf=kopf,
    )
    assert [(f.blatt, f.zeile, f.spalte) for f in erg.fehler] == [("Konferenzen", 3, "Mindestanzahl")]


def test_optionale_werte_aus_dem_blatt(tmp_path):
    kopf = ["Konferenz", "Tag", "Start", "Ende", "Hauptkonferenz", "Wichtigkeit", "Mindestanzahl"]
    erg = lese(tmp_path, ZEILEN, ("Deutsch", *MI, None, 5, 2), ("Sport", *MI2, "Deutsch", 2, 1), kopf=kopf)
    t = {t.name: t for t in erg.themen}
    assert (t["Deutsch"].wichtigkeit, t["Deutsch"].mindestanzahl) == (5, 2)
    assert t["Sport"].haupt == "Deutsch" and t["Sport"].dauer_min == 60


def test_gleichzeitige_konferenzen_rangfolge_entscheidet(tmp_path):
    s = sitzung(tmp_path, ZEILEN, ("Deutsch", *MI), ("Sport", *MI))
    assert {r.name: r.typ for r in s.regeln()}["schwerpunkte"] == "weich"
    plan = s.planen()
    assert plan.status in ("optimal", "machbar")
    assert "Anna Eins" in teilnehmer(plan, "Deutsch") and "Anna Eins" not in teilnehmer(plan, "Sport")
    assert [v.text for v in plan.verstoesse] == ["Anna Eins nimmt nicht an Schwerpunkt Sport teil"]


def test_rangfolge_bei_drei_konferenzen(tmp_path):
    zeilen = [
        ["Eins", "Anna", "", "Deutsch", "Sport", "Mathematik", None],
        ["Zwei", "Ben", "", "Deutsch", None, None, "Deutsch"],
        ["Drei", "Cem", "", "Sport", None, None, "Sport"],
        ["Vier", "Dora", "", "Mathematik", None, None, "Mathematik"],
    ]
    s = sitzung(tmp_path, zeilen, ("Deutsch", *MI), ("Sport", *MI), ("Mathematik", *MI))
    plan = s.planen()
    assert "Anna Eins" in teilnehmer(plan, "Deutsch")
    assert sum("Anna Eins" in t.teilnehmer for t in plan.termine) == 1 and len(plan.verstoesse) == 2


def test_teilueberlappung_zaehlt(tmp_path):
    s = sitzung(tmp_path, ZEILEN, ("Deutsch", "Mi", "14:00", "15:30"), ("Sport", "Mi", "15:00", "16:00"))
    plan = s.planen()
    assert sum("Anna Eins" in t.teilnehmer for t in plan.termine) == 1 and len(plan.verstoesse) == 1


def test_getrennte_konferenzen_ohne_konflikt(tmp_path):
    s = sitzung(tmp_path, ZEILEN, ("Deutsch", *MI), ("Sport", *MI2))
    plan = s.planen()
    assert "Anna Eins" in teilnehmer(plan, "Deutsch") and "Anna Eins" in teilnehmer(plan, "Sport")
    assert not plan.verstoesse


def test_konferenz_mit_zwei_terminen(tmp_path):
    zeilen = [
        ["Eins", "Anna", "", "Deutsch", None, None, None],
        ["Zwei", "Ben", "", "Deutsch", None, None, "Deutsch"],
        ["Drei", "Cem", "", "Deutsch", None, None, None],
    ]
    s = sitzung(tmp_path, zeilen, ("Deutsch", *MI), ("Deutsch", *MI2))
    plan = s.planen()
    assert plan.status in ("optimal", "machbar") and len(plan.termine) == 2
    assert all("Ben Zwei" in t.teilnehmer for t in plan.termine)  # Vorsitz in beiden
    assert (
        teilnehmer(plan, "Deutsch").count("Anna Eins") == 1
        and teilnehmer(plan, "Deutsch").count("Cem Drei") == 1
    )


def test_solver_verschiebt_feste_termine_nicht(tmp_path):
    s = sitzung(tmp_path, ZEILEN, ("Deutsch", "Mi", "16:00", "17:00"), ("Sport", *MI))
    assert s.planen().status in ("optimal", "machbar")
    zeilen = {z[3]: (z[1], z[2]) for z in s.tabellen()["Termine"]}
    assert zeilen["Deutsch"] == (time(16), time(17)) and zeilen["Sport"] == (time(14), time(15))


def test_vorsitz_in_zwei_gleichzeitigen_konferenzen(tmp_path):
    zeilen = [
        ["Eins", "Anna", "", "Deutsch", "Sport", None, "Deutsch"],
        ["Zwei", "Ben", "", "Sport", None, None, "Sport"],
    ]
    # Die Liste erlaubt nur einen Vorsitz je Zeile; Anna wird deshalb im Modell auch Vorsitz von Sport
    s = sitzung(tmp_path, zeilen, ("Deutsch", *MI), ("Sport", *MI))
    from dataclasses import replace

    s.themen = [replace(t, vorsitz="Anna Eins") if t.name == "Sport" else t for t in s.themen]
    plan = s.planen()
    text = " ".join(plan.begruendung)
    assert plan.status == "kein_plan" and "Thema Deutsch" in text and "Thema Sport" in text


def test_standardregel_fuer_lehrerlisten_mit_terminen(tmp_path):
    s = Sitzung()
    s.lade(schreibe(tmp_path / "k.xlsx", lehrer_mit_terminen(ZEILEN, ("Deutsch", *MI), ("Sport", *MI))))
    # Lehrerlisten verteilen sich standardmäßig (SPEC-0012): „alle besuchen“ ist aus
    assert {r.name: r.typ for r in s.regeln()}["schwerpunkte"] == "aus"
    s.setze_profil("festtermine")
    assert {r.name: r.typ for r in s.regeln()}["schwerpunkte"] == "weich"
    s.setze_profil("auto")
    s.lade(
        schreibe(
            tmp_path / "ohne.xlsx",
            {
                "Personen": PERSONEN,
                "Themen": THEMEN,
                "Zeitslots": [["Tag", "Start", "Ende"], ["Mo", "08:00", "08:45"]],
            },
        )
    )
    assert {r.name: r.typ for r in s.regeln()}[
        "schwerpunkte"
    ] == "hart"  # Profil folgt der Datei (Minimalformat)


def test_selbst_gesetzte_regel_bleibt(tmp_path):
    s = Sitzung()
    s.setze_regel("schwerpunkte", "hart", 1.0)
    s.lade(schreibe(tmp_path / "k.xlsx", lehrer_mit_terminen(ZEILEN, ("Deutsch", *MI), ("Sport", *MI))))
    assert {r.name: r.typ for r in s.regeln()}["schwerpunkte"] == "hart"
    assert s.planen().status == "kein_plan"  # hart und gleichzeitig: nicht erfüllbar


def test_selbst_gesetzte_regel_wird_gemerkt_und_standard_nicht(tmp_path):
    from konfiplaner.einstellungen import YamlEinstellungen

    pfad = tmp_path / "e.yaml"
    s = Sitzung(YamlEinstellungen(pfad))
    s.lade(schreibe(tmp_path / "k.xlsx", lehrer_mit_terminen(ZEILEN, ("Deutsch", *MI), ("Sport", *MI))))
    assert (
        "schwerpunkte" not in YamlEinstellungen(pfad).laden()["regeln"]
    )  # nur Profil, nicht von Hand gesetzt
    s.setze_regel("schwerpunkte", "weich", 3.0)
    assert YamlEinstellungen(pfad).laden()["regeln"]["schwerpunkte"]["gewicht"] == 3.0


def test_unterkonferenz_vor_hauptkonferenz_ist_nur_eine_warnung(tmp_path):
    kopf = ["Konferenz", "Tag", "Start", "Ende", "Hauptkonferenz"]
    erg = lese(tmp_path, ZEILEN, ("Deutsch", *MI2, None), ("Sport", *MI, "Deutsch"), kopf=kopf)
    assert not erg.fehler
    assert [(w.blatt, w.zeile, w.spalte) for w in erg.warnungen if "Haupt" in w.spalte] == [
        ("Konferenzen", 3, "Hauptkonferenz")
    ]


def test_unterkonferenz_nach_hauptkonferenz_ohne_warnung(tmp_path):
    kopf = ["Konferenz", "Tag", "Start", "Ende", "Hauptkonferenz"]
    erg = lese(tmp_path, ZEILEN, ("Deutsch", *MI, None), ("Sport", *MI2, "Deutsch"), kopf=kopf)
    assert not erg.meldungen or all("Haupt" not in w.spalte for w in erg.warnungen)


def test_unbekannte_hauptkonferenz_ist_ein_fehler(tmp_path):
    kopf = ["Konferenz", "Tag", "Start", "Ende", "Hauptkonferenz"]
    erg = lese(tmp_path, ZEILEN, ("Deutsch", *MI, "Latein"), ("Sport", *MI2, None), kopf=kopf)
    assert any(f.spalte == "Hauptkonferenz" for f in erg.fehler)


def test_blatt_zeitslots_wird_ignoriert(tmp_path):
    blaetter = lehrer_mit_terminen(ZEILEN, ("Deutsch", *MI), ("Sport", *MI2))
    blaetter["Zeitslots"] = [["Tag", "Start", "Ende"], ["Mo", "08:00", "08:45"]]
    erg = importiere(schreibe(tmp_path / "z.xlsx", blaetter))
    assert not erg.fehler and any("ignoriert" in w.text for w in erg.warnungen)
    assert all(z.tag == "Mi" for z in erg.zeitslots)


def test_unbekannter_konferenzname_im_blatt_ist_ein_fehler(tmp_path):
    erg = lese(tmp_path, ZEILEN, ("Deutsch", *MI), ("Sport", *MI2), ("Latein", *MI))
    assert [(f.zeile, f.spalte) for f in erg.fehler] == [(4, "Konferenz")]


def test_bestaetigte_konferenz_ohne_mitglieder_wird_ignoriert(tmp_path):
    n = Normalisierung()
    n.bestaetige("Latein", ["Latein"])
    erg = lese(tmp_path, ZEILEN, ("Deutsch", *MI), ("Sport", *MI2), ("Latein", *MI), norm=n)
    assert not erg.fehler and any("keine Mitglieder" in w.text for w in erg.warnungen)
    assert "Latein" not in [t.name for t in erg.themen]


def test_minimalformat_mit_blatt_konferenzen(tmp_path):
    blaetter = {
        "Personen": PERSONEN,
        "Themen": THEMEN,
        "Konferenzen": [
            ["Konferenz", "Tag", "Start", "Ende"],
            ["Taufe", "Mo", "08:00", "09:00"],
            ["Abendmahl", "Mo", "09:00", "10:30"],
        ],
    }
    erg = importiere(schreibe(tmp_path / "m.xlsx", blaetter))
    assert not erg.fehler and erg.festtermine and erg.modell is not None
    s = Sitzung()
    s.lade(tmp_path / "m.xlsx")
    assert s.planen().status in ("optimal", "machbar")


def test_konflikte_im_ergebnis_und_export(tmp_path):
    s = sitzung(tmp_path, ZEILEN, ("Deutsch", *MI), ("Sport", *MI))
    s.planen()
    hinweise = s.tabellen()["Hinweise"]
    assert ["Verstoß", "schwerpunkte: Anna Eins nimmt nicht an Schwerpunkt Sport teil"] in hinweise
    ziel = tmp_path / "plan.xlsx"
    s.exportiere(ziel)
    blatt = [[z.value for z in r] for r in load_workbook(ziel)["Hinweise"].iter_rows(min_row=2)]
    assert ["Verstoß", "schwerpunkte: Anna Eins nimmt nicht an Schwerpunkt Sport teil"] in blatt


def test_zeitraster_formular_ist_bei_festterminen_gesperrt(tmp_path):
    import pytest

    from konfiplaner.modell import ModellFehler

    s = sitzung(tmp_path, ZEILEN, ("Deutsch", *MI), ("Sport", *MI2))
    with pytest.raises(ModellFehler, match="Blatt Konferenzen"):
        s.erzeuge_zeitslots("Mi", time(9), time(12), 45)


def test_dateien_ohne_blatt_konferenzen_verhalten_sich_wie_bisher(tmp_path):
    from tests.xlsx_helfer import lehrerliste

    erg = importiere(schreibe(tmp_path / "l.xlsx", lehrerliste(LEHRER)))
    assert not erg.festtermine and erg.zeitslots and all(not t.festtermine for t in erg.themen)


def test_planung_der_realistischen_groesse(tmp_path):
    import random
    import time as t

    rnd = random.Random(3)
    konferenzen = [f"K{i}" for i in range(26)]
    zeilen = [
        [f"N{i}", f"V{i}", "", *rnd.sample(konferenzen, 2), None, konferenzen[i] if i < 26 else None]
        for i in range(100)
    ]
    fenster = [("Mi", "14:00", "15:30"), ("Mi", "15:00", "16:30"), ("Do", "14:00", "15:30")]
    s = sitzung(tmp_path, zeilen, *[(k, *rnd.choice(fenster)) for k in konferenzen])
    s.zeitlimit_s = 30
    t0 = t.perf_counter()
    plan = s.planen()
    assert plan.status in ("optimal", "machbar") and t.perf_counter() - t0 < 30
