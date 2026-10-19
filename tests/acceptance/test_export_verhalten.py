"""TST-0009 – Export-Verhalten (CON-0009). Die Testnamen folgen den Szenarien der .feature-Datei."""

from datetime import time

import pytest
from openpyxl import load_workbook

from konfiplaner.excel_export import ExportFehler, ZielDateiExistiertFehler, exportiere
from konfiplaner.modell import Plan, Termin, Thema, Verstoss
from tests.helfer import modell


def zeilen(pfad, blatt):
    return [[z.value for z in r] for r in load_workbook(pfad)[blatt].iter_rows(min_row=2)]


@pytest.fixture
def m():
    return modell(
        [Thema("Taufe", mindestanzahl=2), Thema("Abendmahl", mindestanzahl=3)],
        [("Anna", "Taufe", "Abendmahl"), ("Ben", "Taufe", "Abendmahl"), ("Cem", "Abendmahl", "Taufe")],
        6,
    )


@pytest.fixture
def plan():
    return Plan(
        "optimal",
        [Termin("Taufe", 1, 2, ("Anna", "Ben")), Termin("Abendmahl", 3, 1, ("Anna", "Ben"))],
        seed=1,
    )


def test_plan_exportieren(tmp_path, m, plan):
    ziel = tmp_path / "plan.xlsx"
    exportiere(plan, m, ziel)
    assert load_workbook(ziel).sheetnames == ["Termine", "Personen", "Themen", "Hinweise"]
    personen = zeilen(ziel, "Personen")
    assert [r[4] for r in personen if r[0] == "Anna"] == ["Taufe", "Abendmahl"]  # zeitliche Reihenfolge
    assert {r[0] for r in personen} == {"Anna", "Ben"}


def test_termine_mit_zeit_und_teilnehmenden(tmp_path, m, plan):
    ziel = tmp_path / "plan.xlsx"
    exportiere(plan, m, ziel)
    erste = zeilen(ziel, "Termine")[0]
    assert erste == [
        "Mo",
        time(8, 45),
        time(10, 15),
        "Taufe",
        None,
        2,
        "Anna, Ben",
    ]  # leerer Vorsitz = leere Zelle


def test_themenauslastung(tmp_path, m, plan):
    ziel = tmp_path / "plan.xlsx"
    exportiere(plan, m, ziel)
    zeile = {r[0]: r for r in zeilen(ziel, "Themen")}
    assert zeile["Taufe"][4] == "ja" and zeile["Abendmahl"][4] == "nein"
    assert zeile["Abendmahl"][3] == 2


def test_kein_plan(tmp_path, m):
    ziel = tmp_path / "plan.xlsx"
    exportiere(Plan("kein_plan", begruendung=["Regel mindestanzahl: Thema A"]), m, ziel)
    assert zeilen(ziel, "Termine") == []
    assert ["Begründung", "Regel mindestanzahl: Thema A"] in zeilen(ziel, "Hinweise")
    assert all(r[4] == "nein" for r in zeilen(ziel, "Themen"))


def test_plan_nicht_optimal_und_verstoesse_im_hinweisblatt(tmp_path, m):
    ziel = tmp_path / "plan.xlsx"
    plan = Plan(
        "machbar",
        [Termin("Taufe", 0, 1, ("Anna", "Ben"))],
        [Verstoss("schwerpunkte", "Cem fehlt")],
        warnungen=["Dora hat keine Verfügbarkeit"],
    )
    exportiere(plan, m, ziel)
    hinweise = zeilen(ziel, "Hinweise")
    assert ["Status", "Plan ist nicht optimal"] in hinweise
    assert ["Verstoß", "schwerpunkte: Cem fehlt"] in hinweise
    assert ["Warnung", "Dora hat keine Verfügbarkeit"] in hinweise


def test_zieldatei_existiert(tmp_path, m, plan):
    ziel = tmp_path / "plan.xlsx"
    ziel.write_text("alt")
    with pytest.raises(ZielDateiExistiertFehler):
        exportiere(plan, m, ziel)
    assert ziel.read_text() == "alt"
    exportiere(plan, m, ziel, ueberschreiben=True)
    assert load_workbook(ziel).sheetnames[0] == "Termine"


def test_zieldatei_gesperrt_durch_sperrdatei(tmp_path, m, plan):
    ziel = tmp_path / "plan.xlsx"
    (tmp_path / "~$plan.xlsx").write_text("")
    with pytest.raises(ExportFehler, match="geöffnet"):
        exportiere(plan, m, ziel)


def test_schreibfehler_wird_klar_gemeldet_und_zerstoert_nichts(tmp_path, m, plan, monkeypatch):
    ziel = tmp_path / "plan.xlsx"
    ziel.write_text("alt")

    def gesperrt(*_a, **_k):
        raise PermissionError("gesperrt")

    monkeypatch.setattr("os.replace", gesperrt)
    with pytest.raises(ExportFehler, match="plan.xlsx"):
        exportiere(plan, m, ziel, ueberschreiben=True)
    assert ziel.read_text() == "alt"
    assert [p.name for p in tmp_path.iterdir()] == ["plan.xlsx"]  # keine Temp-Reste


def test_termin_ausserhalb_der_slots(tmp_path, m):
    with pytest.raises(ExportFehler):
        exportiere(Plan("optimal", [Termin("Taufe", 5, 3, ())]), m, tmp_path / "x.xlsx")


def test_ende_zu_ende_vom_solver_zum_excel(tmp_path):
    from konfiplaner.solver import CpSatPlaner
    from tests.helfer import mit_regeln

    m = modell(
        [Thema("A", mindestanzahl=2), Thema("B", mindestanzahl=2)], [(f"P{i}", "A", "B") for i in range(4)], 4
    )
    plan = CpSatPlaner().planen(m, mit_regeln())
    ziel = tmp_path / "plan.xlsx"
    exportiere(plan, m, ziel)
    assert {r[3] for r in zeilen(ziel, "Termine")} == {"A", "B"}
    assert len(zeilen(ziel, "Personen")) == 8
