"""TST-0013 – Vorsitz in Modell, Import, Export und Sitzung (CON-0011)."""

import json
from pathlib import Path

import pytest
from jsonschema import Draft202012Validator
from openpyxl import load_workbook

from konfiplaner.anwendung import Sitzung
from konfiplaner.excel_export import SPALTEN, exportiere
from konfiplaner.excel_import import importiere
from konfiplaner.modell import ModellFehler, Plan, Termin, Thema
from tests.helfer import modell
from tests.xlsx_helfer import PERSONEN, SLOTS, schreibe

THEMEN_V = [
    ["Name", "Wichtigkeit", "Mindestanzahl", "Dauer", "Vorsitz"],
    ["Taufe", 3, 2, 45, "Anna"],
    ["Abendmahl", 5, 2, 90, None],
]
IMPORT_SCHEMA = json.loads(Path(".sdd/contracts/data/importmodell-und-fehlerliste.schema.json").read_text())
EXPORT_SCHEMA = json.loads(Path(".sdd/contracts/data/export-blattstruktur.schema.json").read_text())


def lese(tmp_path, themen, **kw):
    return importiere(
        schreibe(tmp_path / "t.xlsx", {"Personen": PERSONEN, "Themen": themen, "Zeitslots": SLOTS}), **kw
    )


# ── Modell ──────────────────────────────────────────────────────────────────
def test_vorsitz_muss_eine_person_sein():
    with pytest.raises(ModellFehler, match="Vorsitz"):
        modell([Thema("A", vorsitz="Niemand")], [("P", "A", "B")][:0])


def test_personen_zu_und_themen_von_beruecksichtigen_den_vorsitz():
    m = modell([Thema("A", vorsitz="Z"), Thema("B"), Thema("C")], [("P", "B", "C"), ("Z", "B", "C")])
    assert [p.name for p in m.personen_zu("A")] == ["Z"]
    assert m.themen_von("Z") == ["A", "B", "C"] and m.themen_von("P") == ["B", "C"]


def test_vorsitz_ist_optional():
    assert Thema("A").vorsitz is None


# ── Import ──────────────────────────────────────────────────────────────────
def test_import_uebernimmt_den_vorsitz(tmp_path):
    erg = lese(tmp_path, THEMEN_V)
    assert not erg.meldungen
    assert [t.vorsitz for t in erg.themen] == ["Anna", None]
    assert [t.vorsitz for t in erg.modell.themen] == ["Anna", None]


@pytest.mark.parametrize("kopf", ["vorsitz", "Vorsitzender", " LEITUNG "])
def test_import_erkennt_ueberschrift_varianten(tmp_path, kopf):
    erg = lese(tmp_path, [[*THEMEN_V[0][:4], kopf], *THEMEN_V[1:]])
    assert erg.themen[0].vorsitz == "Anna"


def test_import_unbekannter_vorsitz_ist_ein_fehler_mit_stelle(tmp_path):
    themen = [THEMEN_V[0], ["Taufe", 3, 2, 45, "Unbekannt"], THEMEN_V[2]]
    erg = lese(tmp_path, themen)
    assert [(f.blatt, f.zeile, f.spalte) for f in erg.fehler] == [("Themen", 2, "Vorsitz")]
    assert "Unbekannt" in erg.fehler[0].text and erg.modell is None


def test_import_ohne_spalte_vorsitz_wie_bisher(tmp_path):
    themen = [["Name", "Wichtigkeit", "Mindestanzahl", "Dauer"], ["Taufe", 3, 2, 45], ["Abendmahl", 5, 2, 90]]
    erg = lese(tmp_path, themen)
    assert not erg.meldungen and all(t.vorsitz is None for t in erg.themen)


def test_import_leere_vorsitz_zelle_bedeutet_keinen_vorsitz(tmp_path):
    themen = [THEMEN_V[0], ["Taufe", 3, 2, 45, "   "], THEMEN_V[2]]
    assert lese(tmp_path, themen).themen[0].vorsitz is None


def test_importergebnis_entspricht_dem_schema_mit_vorsitz(tmp_path):
    d = lese(tmp_path, THEMEN_V).to_dict()
    assert d["modell"]["themen"][0]["vorsitz"] == "Anna" and "vorsitz" not in d["modell"]["themen"][1]
    assert not list(Draft202012Validator(IMPORT_SCHEMA).iter_errors(d))


# ── Export ──────────────────────────────────────────────────────────────────
def test_export_zeigt_vorsitz_und_rolle(tmp_path):
    m = modell(
        [Thema("Taufe", vorsitz="Anna"), Thema("Gebet")],
        [("Anna", "Gebet", "Taufe"), ("Ben", "Taufe", "Gebet")],
    )
    plan = Plan("optimal", [Termin("Taufe", 0, 1, ("Anna", "Ben")), Termin("Gebet", 1, 1, ("Anna", "Ben"))])
    ziel = tmp_path / "plan.xlsx"
    exportiere(plan, m, ziel)
    mappe = load_workbook(ziel)
    termine = [[z.value for z in r] for r in mappe["Termine"].iter_rows(min_row=2)]
    assert [(r[3], r[4]) for r in termine] == [("Taufe", "Anna"), ("Gebet", None)]
    personen = [[z.value for z in r] for r in mappe["Personen"].iter_rows(min_row=2)]
    rollen = {(r[0], r[4]): r[5] for r in personen}
    assert rollen[("Anna", "Taufe")] == "Vorsitz" and rollen[("Ben", "Taufe")] == "Teilnahme"
    assert rollen[("Anna", "Gebet")] == "Teilnahme"


def test_export_spalten_entsprechen_dem_schema():
    for blatt, spalten in SPALTEN.items():
        assert EXPORT_SCHEMA["properties"][blatt]["const"] == spalten
    assert "Vorsitz" in SPALTEN["Termine"] and SPALTEN["Personen"][-1] == "Rolle"


# ── Sitzung ─────────────────────────────────────────────────────────────────
def test_sitzung_setzt_und_entfernt_vorsitz(tmp_path):
    s = Sitzung()
    s.lade(schreibe(tmp_path / "t.xlsx", {"Personen": PERSONEN}))
    assert s.personen_namen() == ["Anna", "Ben"]
    s.setze_thema(0, vorsitz="Ben")
    assert s.themen[0].vorsitz == "Ben"
    s.setze_thema(0, vorsitz="")
    assert s.themen[0].vorsitz is None


def test_sitzung_lehnt_unbekannten_vorsitz_ab(tmp_path):
    s = Sitzung()
    s.lade(schreibe(tmp_path / "t.xlsx", {"Personen": PERSONEN}))
    with pytest.raises(ModellFehler, match="Unbekannte Person"):
        s.setze_thema(0, vorsitz="Niemand")
    assert s.themen[0].vorsitz is None


def test_sitzung_plant_mit_vorsitz_aus_der_datei(tmp_path):
    personen = [
        PERSONEN[0],
        ["Anna", "Taufe", "Abendmahl"],
        ["Ben", "Abendmahl", "Taufe"],
        ["Cem", "Gebet", "Gebet2"],
    ]
    themen = [
        ["Name", "Wichtigkeit", "Mindestanzahl", "Dauer", "Vorsitz"],
        ["Taufe", 3, 2, 45, "Cem"],
        ["Abendmahl", 3, 2, 45, None],
        ["Gebet", 3, 1, 45, None],
        ["Gebet2", 3, 1, 45, None],
    ]
    s = Sitzung()
    s.lade(schreibe(tmp_path / "t.xlsx", {"Personen": personen, "Themen": themen, "Zeitslots": SLOTS}))
    plan = s.planen()
    assert plan.status in ("optimal", "machbar")
    assert all("Cem" in t.teilnehmer for t in plan.termine if t.thema == "Taufe")
