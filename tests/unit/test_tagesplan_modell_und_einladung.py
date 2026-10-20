"""TST-0029 – Tagesplan-Modell, Prüfung, Persistenz und Word-Inhalt (CON-0019)."""

from datetime import date, time

import docx
import pytest
from docx.enum.section import WD_ORIENT

from konfiplaner.anwendung import Sitzung
from konfiplaner.einladung import exportiere_einladung
from konfiplaner.einstellungen import YamlEinstellungen
from konfiplaner.excel_export import ExportFehler, ZielDateiExistiertFehler
from konfiplaner.tagesplan import Block, Tagesplan, standard_bloecke
from tests.unit.test_auswahl_in_der_sitzung import MI, ZEILEN
from tests.xlsx_helfer import lehrer_mit_terminen, schreibe

PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360f8cfc00000"
    "03010100c9fe92ef0000000049454e44ae426082"
)


def test_standardablauf_hat_die_erwarteten_bloecke():
    titel = [b.titel for b in Tagesplan().bloecke]
    assert titel[0] == "Stehkaffee im Foyer" and titel[-1] == "Ende der Veranstaltung"
    assert any(b.art == "konferenzen" for b in standard_bloecke()) and len(titel) == 9


@pytest.mark.parametrize(
    ("block", "text"),
    [
        (Block(time(8, 30), time(8, 40)), "8.30 – 8.40 Uhr"),
        (Block(time(8, 0), None, "ab"), "ab 8.00 Uhr"),
        (Block(time(15, 0), None, "um"), "15.00 Uhr"),
        (Block(time(13, 5), time(14, 30)), "13.05 – 14.30 Uhr"),
    ],
)
def test_zeitformat(block, text):
    assert block.zeit_text == text


def test_konferenzblock_traegt_konferenzen_vor_dem_eigenen_inhalt_ein():
    plan = Tagesplan()
    block = Block(time(9), time(10), "bereich", "FK", "Material sichten", "", "konferenzen")
    assert plan.block_inhalt(block, ["Mathematik", "Deutsch"]) == "Mathematik, Deutsch\nMaterial sichten"
    assert plan.block_inhalt(block, []) == "Material sichten"
    frei = Block(time(9), time(10), "bereich", "x", "Text")
    assert plan.block_inhalt(frei, ["Mathematik"]) == "Text"


def test_pruefung_fehler_und_warnung():
    ende_vor_beginn = Tagesplan(bloecke=(Block(time(10), time(9)),))
    assert [h.fehler for h in ende_vor_beginn.pruefe()] == [True]
    gleich = Tagesplan(bloecke=(Block(time(10), time(10)),))
    assert [h.fehler for h in gleich.pruefe()] == [True]
    ueberlappung = Tagesplan(bloecke=(Block(time(9), time(10)), Block(time(9, 30), time(10, 30))))
    assert [h.fehler for h in ueberlappung.pruefe()] == [False]
    anschliessend = Tagesplan(bloecke=(Block(time(9), time(10)), Block(time(10), time(11))))
    assert anschliessend.pruefe() == [] and Tagesplan().pruefe() == []


def test_dict_runde_und_fehlertoleranz():
    plan = Tagesplan(datum=date(2026, 11, 9), thema="T", ziele=("a", "b"), vorbereitung="V")
    assert Tagesplan.aus_dict(plan.als_dict()) == plan
    assert Tagesplan.aus_dict("kaputt") == Tagesplan()
    kaputt = Tagesplan.aus_dict({"datum": "x", "bloecke": [{"beginn": "??"}, {"beginn": "9:00", "art": "?"}]})
    assert kaputt.datum is None and len(kaputt.bloecke) == 1 and kaputt.bloecke[0].art == "frei"


def exportiert(tmp_path, plan, konferenzen=(), **kw):
    ziel = tmp_path / "e.docx"
    exportiere_einladung(plan, list(konferenzen), ziel, True, **kw)
    return docx.Document(str(ziel))


def test_dokument_enthaelt_kopf_und_ablauftabelle(tmp_path):
    plan = Tagesplan(
        datum=date(2026, 11, 9), titel="Fortbildung", thema="Methoden", ziele=("Ziel A", "Ziel B")
    )
    d = exportiert(tmp_path, plan, ["Mathematik", "Deutsch"])
    texte = [p.text for p in d.paragraphs]
    assert texte[0] == "Montag, 09.11.2026: Fortbildung" and "THEMA: Methoden" in texte
    assert "Ablaufplan 09.11.2026" in texte and "Ziel A" in texte and "Ziel B" in texte
    assert [c.text for c in d.tables[0].rows[0].cells] == ["Zeiten", "Inhalt", "Wer/Wie?"]
    konferenz_zeile = next(r for r in d.tables[0].rows if "Fachkonferenzen" in r.cells[1].text)
    assert "Mathematik, Deutsch" in konferenz_zeile.cells[1].text
    assert konferenz_zeile.cells[1].paragraphs[0].runs[0].bold  # Blocktitel fett
    abschnitt = d.sections[0]
    assert abschnitt.orientation == WD_ORIENT.LANDSCAPE and abschnitt.page_width > abschnitt.page_height


def test_leere_abschnitte_entfallen_und_ohne_datum(tmp_path):
    d = exportiert(tmp_path, Tagesplan(titel="X"))
    texte = [p.text for p in d.paragraphs]
    assert not any(
        t.startswith(("Vorbereitung", "Arbeitsformen", "Verschriftlichung", "THEMA")) for t in texte
    )
    assert texte[0] == "X" and "Ablaufplan" in texte


def test_logo_wird_eingefuegt_oder_uebergangen(tmp_path):
    logo = tmp_path / "logo.png"
    logo.write_bytes(PNG)
    assert len(exportiert(tmp_path, Tagesplan(), logo=logo).inline_shapes) == 1
    assert len(exportiert(tmp_path, Tagesplan(), logo=tmp_path / "fehlt.png").inline_shapes) == 0
    kaputt = tmp_path / "kaputt.png"
    kaputt.write_bytes(b"kein bild")
    assert len(exportiert(tmp_path, Tagesplan(), logo=kaputt).inline_shapes) == 0


def test_fehler_im_plan_schreiben_keine_datei(tmp_path):
    ziel = tmp_path / "e.docx"
    with pytest.raises(ExportFehler, match="Ende"):
        exportiere_einladung(Tagesplan(bloecke=(Block(time(10), time(9), "bereich", "X"),)), [], ziel)
    assert not ziel.exists()


def test_vorhandene_datei_nur_nach_bestaetigung(tmp_path):
    ziel = tmp_path / "e.docx"
    ziel.write_text("alt")
    with pytest.raises(ZielDateiExistiertFehler):
        exportiere_einladung(Tagesplan(), [], ziel)
    assert ziel.read_text() == "alt"
    exportiere_einladung(Tagesplan(), [], ziel, True)
    assert docx.Document(str(ziel)).tables


def sitzung(tmp_path):
    datei = schreibe(
        tmp_path / "k.xlsx",
        lehrer_mit_terminen(ZEILEN, ("Deutsch", *MI), ("Sport", *MI), ("Kunst", "Mi", "15:00", "16:00")),
    )
    s = Sitzung(YamlEinstellungen(tmp_path / "einstellungen.yaml"))
    s.lade(datei)
    return s


def test_sitzung_exportiert_gewaehlte_konferenzen_und_merkt_sich_den_plan(tmp_path):
    s = sitzung(tmp_path)
    alle = s.ausgewaehlte_konferenzen()
    s.abgewaehlt.add(alle[-1])
    s.setze_tagesplan(Tagesplan(titel="Gespeichert"))
    s.exportiere_einladung(tmp_path / "e.docx")
    d = docx.Document(str(tmp_path / "e.docx"))
    zeile = next(r for r in d.tables[0].rows if "Fachkonferenzen" in r.cells[1].text)
    assert alle[0] in zeile.cells[1].text and alle[-1] not in zeile.cells[1].text
    neu = Sitzung(YamlEinstellungen(tmp_path / "einstellungen.yaml"))
    assert neu.tagesplan.titel == "Gespeichert" and neu.letzter_exportordner == tmp_path
    s.setze_logo(tmp_path / "l.png")
    assert Sitzung(YamlEinstellungen(tmp_path / "einstellungen.yaml")).logo == tmp_path / "l.png"
