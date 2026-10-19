"""TST-0008 – Export-Blattstruktur (CON-0008)."""

import json
from pathlib import Path

from jsonschema import Draft202012Validator
from openpyxl import load_workbook

from konfiplaner.excel_export import SPALTEN, exportiere
from konfiplaner.modell import Plan, Termin, Thema
from tests.helfer import modell

SCHEMA = json.loads(Path(".sdd/contracts/data/export-blattstruktur.schema.json").read_text())


def _struktur(pfad) -> dict:
    mappe = load_workbook(pfad)
    return {blatt.title: [z.value for z in blatt[1]] for blatt in mappe}


def test_export_entspricht_dem_schema(tmp_path):
    m = modell([Thema("A")], [("P", "A", "A2")][:0])
    ziel = tmp_path / "plan.xlsx"
    exportiere(Plan("optimal", [Termin("A", 0, 1, ())]), m, ziel)
    struktur = _struktur(ziel)
    assert not list(Draft202012Validator(SCHEMA).iter_errors(struktur))


def test_blaetter_in_vertraglicher_reihenfolge(tmp_path):
    ziel = tmp_path / "plan.xlsx"
    exportiere(Plan("kein_plan", begruendung=["x"]), modell([Thema("A")], []), ziel)
    assert load_workbook(ziel).sheetnames == ["Termine", "Personen", "Themen", "Hinweise"]


def test_spalten_konstante_entspricht_dem_schema():
    for blatt, spalten in SPALTEN.items():
        assert SCHEMA["properties"][blatt]["const"] == spalten
