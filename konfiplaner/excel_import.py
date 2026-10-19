"""Excel-Import von Personen, Themen, Zeitslots und Verfügbarkeiten (SPEC-0001).

Der Import kennt weder Solver noch GUI: Er erzeugt nur Objekte des Domänenmodells.
"""

from __future__ import annotations

import dataclasses
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from konfiplaner.lehrerliste import ist_lehrerliste, lese_lehrerliste
from konfiplaner.modell import Modell, ModellFehler, Person, Thema, Zeitslot
from konfiplaner.normalisierung import SICHER, VORSCHLAG, Aufloesung, Normalisierung
from konfiplaner.zeitraster import raster_aus_fenstern

STANDARD_WICHTIGKEIT = 3
STANDARD_MINDESTANZAHL = 1
STANDARD_DAUER_MIN = 45

BLAETTER = {
    "Personen": ("name", "thema_1", "thema_2"),
    "Themen": ("name", "wichtigkeit", "mindestanzahl", "dauer"),
    "Zeitslots": ("tag", "start", "ende"),
    "Verfügbarkeiten": ("person", "tag", "start"),
    "Konferenzen": ("konferenz", "tag", "start", "ende"),
}
PFLICHTBLAETTER = ("Personen",)
# Spalten, die fehlen dürfen (kein Fehler, kein Wert)
OPTIONALE_SPALTEN = {
    "Themen": ("vorsitz",),
    "Konferenzen": ("hauptkonferenz", "wichtigkeit", "mindestanzahl"),
}

# Kanonischer Spaltenname -> akzeptierte Überschriften (normalisiert, siehe `_norm`)
STANDARD_ALIASSE: dict[str, tuple[str, ...]] = {
    "name": ("name", "vorname nachname", "konfirmand", "person"),
    "thema_1": ("thema 1", "thema1", "schwerpunkt 1", "schwerpunkt1", "erstes thema"),
    "thema_2": ("thema 2", "thema2", "schwerpunkt 2", "schwerpunkt2", "zweites thema"),
    "wichtigkeit": ("wichtigkeit", "prioritaet", "priorität"),
    "mindestanzahl": ("mindestanzahl", "mindestteilnehmer", "min"),
    "dauer": ("dauer", "dauer min", "dauer in minuten", "minuten"),
    "vorsitz": ("vorsitz", "vorsitzender", "vorsitzende", "leitung"),
    "konferenz": ("konferenz", "fachkonferenz", "thema", "name"),
    "hauptkonferenz": ("hauptkonferenz", "haupt", "übergeordnet"),
    "tag": ("tag", "datum"),
    "start": ("start", "beginn", "von"),
    "ende": ("ende", "bis"),
    "person": ("person", "name"),
}


class ImportDateiFehler(Exception):
    """Die Datei ist keine lesbare Excel-Datei (beschädigt, falsches Format, passwortgeschützt)."""


@dataclass(frozen=True)
class Meldung:
    schwere: str  # "fehler" | "warnung"
    blatt: str
    zeile: int
    spalte: str
    text: str


@dataclass
class Importergebnis:
    """Ergebnis eines Imports. `modell` gibt es nur ohne Fehler und mit vorhandenen Zeitslots."""

    meldungen: list[Meldung] = field(default_factory=list)
    personen: list[Person] = field(default_factory=list)
    themen: list[Thema] = field(default_factory=list)
    themen_aus_blatt: bool = False
    zeitslots: list[Zeitslot] = field(default_factory=list)
    verfuegbarkeit: dict[str, set[int]] = field(default_factory=dict)
    quelle: dict[tuple[str, int], str] = field(default_factory=dict)
    # Lehrerliste (SPEC-0007)
    lehrerliste: bool = False
    aufloesungen: list[Aufloesung] = field(default_factory=list)
    ohne_konferenz: list[str] = field(default_factory=list)
    festtermine: bool = False  # Blatt Konferenzen vorhanden (SPEC-0008)

    @property
    def fehler(self) -> list[Meldung]:
        return [m for m in self.meldungen if m.schwere == "fehler"]

    @property
    def warnungen(self) -> list[Meldung]:
        return [m for m in self.meldungen if m.schwere == "warnung"]

    @property
    def modell(self) -> Modell | None:
        if self.fehler or not self.zeitslots:
            return None
        return self.baue_modell(self.themen, self.zeitslots, self.verfuegbarkeit)

    def to_dict(self) -> dict:
        """Darstellung gemäß CON-0001 (Schema `importmodell-und-fehlerliste`)."""
        d: dict = {
            "meldungen": [
                {"schwere": m.schwere, "blatt": m.blatt, "zeile": m.zeile, "spalte": m.spalte, "text": m.text}
                for m in self.meldungen
            ]
        }
        modell = self.modell
        if modell is not None:
            d["modell"] = {
                "personen": [{"name": p.name, "themen": list(p.themen)} for p in modell.personen],
                "themen": [
                    {
                        "name": t.name,
                        "wichtigkeit": t.wichtigkeit,
                        "mindestanzahl": t.mindestanzahl,
                        "dauer_min": t.dauer_min,
                        **({"vorsitz": t.vorsitz} if t.vorsitz else {}),
                    }
                    for t in modell.themen
                ],
                "zeitslots": [
                    {"tag": s.tag, "start": f"{s.start:%H:%M}", "ende": f"{s.ende:%H:%M}"}
                    for s in modell.zeitslots
                ],
                "verfuegbarkeiten": [
                    {"person": p, "slot": str(i)}
                    for p, slots in sorted(modell.verfuegbarkeit.items())
                    for i in sorted(slots)
                ],
            }
        return d

    def baue_modell(
        self,
        themen: list[Thema],
        zeitslots: list[Zeitslot],
        verfuegbarkeit: dict[str, set[int]] | None = None,
        personen: list[Person] | None = None,
    ) -> Modell:
        """Modell aus den importierten Personen (oder `personen`) und den Themen und Slots."""
        return Modell(
            themen, self.personen if personen is None else personen, zeitslots, verfuegbarkeit or {}
        )


def _norm(text: Any) -> str:
    return re.sub(r"[\s_\-.:]+", " ", str(text or "").strip().lower()).strip()


def _leer(wert: Any) -> bool:
    return wert is None or (isinstance(wert, str) and not wert.strip())


def importiere(
    pfad: str | Path,
    zuordnung: dict[str, list[str]] | None = None,
    normalisierung: Normalisierung | None = None,
) -> Importergebnis:
    """Liest eine `.xlsx`-Datei (Minimalformat oder Lehrerliste, SPEC-0001 FR-11).

    `zuordnung` ergänzt die Überschriften-Aliasse je Spalte, `normalisierung` enthält bestätigte
    Konferenz-Zuordnungen der Lehrerliste.
    """
    pfad = Path(pfad)
    try:
        mappe = load_workbook(pfad, data_only=True, read_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as e:
        raise ImportDateiFehler(
            f"{pfad.name} ist keine lesbare Excel-Datei (beschädigt, passwortgeschützt oder falsches Format)"
        ) from e
    try:
        return _Leser(mappe, zuordnung or {}, normalisierung or Normalisierung()).lese()
    finally:
        mappe.close()


def finde_spalte(kopfzeile: list[str], spalte: str) -> int | None:
    """Index der Kopfzeilen-Zelle, die zur Meldungs-Spalte (z. B. "Thema 1") gehört."""
    for i, zelle in enumerate(kopfzeile):  # Überschrift stimmt wörtlich (Lehrerliste)
        if zelle and _norm(zelle) == _norm(spalte):
            return i
    kanon = _norm(spalte).replace(" ", "_")
    aliasse = {_norm(a) for a in STANDARD_ALIASSE.get(kanon, (kanon,))}
    for i, zelle in enumerate(kopfzeile):
        if _norm(zelle) in aliasse:
            return i
    return None


def rohdaten(pfad: str | Path) -> dict[str, list[list[str]]]:
    """Alle Blätter als Textzeilen (inkl. Kopfzeile, Index = Excel-Zeile - 1) für Vorschauen."""
    try:
        mappe = load_workbook(Path(pfad), data_only=True, read_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as e:
        raise ImportDateiFehler(f"{Path(pfad).name} ist keine lesbare Excel-Datei") from e
    try:
        return {
            ws.title: [
                [
                    "" if c is None else (f"{c:%H:%M}" if isinstance(c, (time, datetime)) else str(c))
                    for c in zeile
                ]
                for zeile in ws.iter_rows(values_only=True)
            ]
            for ws in mappe
        }
    finally:
        mappe.close()


class _Leser:
    def __init__(self, mappe, zuordnung: dict[str, list[str]], normalisierung: Normalisierung) -> None:
        self.mappe = mappe
        self.normalisierung = normalisierung
        self.aliasse = {
            k: {_norm(a) for a in (*STANDARD_ALIASSE.get(k, (k,)), *zuordnung.get(k, ()))}
            for k in {*STANDARD_ALIASSE, *zuordnung}
        }
        self.erg = Importergebnis()
        self._themen_zeilen: dict[str, int] = {}
        self._blaetter = {_norm(n): n for n in mappe.sheetnames}
        if "konferenzen" not in self._blaetter and "konferenztermine" in self._blaetter:
            self._blaetter["konferenzen"] = self._blaetter["konferenztermine"]

    # ── Hilfen ──────────────────────────────────────────────────────────────
    def _meld(self, schwere: str, blatt: str, zeile: int, spalte: str, text: str) -> None:
        self.erg.meldungen.append(Meldung(schwere, blatt, zeile, spalte, text))

    def _zeilen(self, blatt: str) -> tuple[dict[str, int], list[tuple[int, list[Any]]]] | None:
        """Spaltenindex je kanonischem Namen und die nicht leeren Datenzeilen (Excel-Zeilennummer, Werte)."""
        echt = self._blaetter.get(_norm(blatt))
        if echt is None:
            return None
        ws = self.mappe[echt]
        zeilen = list(ws.iter_rows(values_only=True))
        if not zeilen:
            self._meld("fehler", blatt, 1, "", "Blatt ist leer")
            return {}, []
        kopf = [_norm(c) for c in zeilen[0]]
        spalten: dict[str, int] = {}
        for kanon in (*BLAETTER[blatt], *OPTIONALE_SPALTEN.get(blatt, ())):
            for i, h in enumerate(kopf):
                if h in self.aliasse[kanon]:
                    spalten[kanon] = i
                    break
            else:
                if kanon in OPTIONALE_SPALTEN.get(blatt, ()):
                    continue
                self._meld(
                    "fehler",
                    blatt,
                    1,
                    kanon.replace("_", " ").capitalize(),
                    f"Spalte {kanon.replace('_', ' ')} fehlt",
                )
        daten = [(i, list(z)) for i, z in enumerate(zeilen[1:], start=2) if not all(_leer(c) for c in z)]
        return spalten, daten

    @staticmethod
    def _wert(zeile: list[Any], spalten: dict[str, int], name: str) -> Any:
        i = spalten.get(name)
        return zeile[i] if i is not None and i < len(zeile) else None

    def _text(self, blatt, nr, zeile, spalten, name, titel) -> str | None:
        w = self._wert(zeile, spalten, name)
        if _leer(w):
            self._meld("fehler", blatt, nr, titel, f"{titel} fehlt")
            return None
        return str(w).strip()

    def _zahl(
        self, blatt, nr, zeile, spalten, name, titel, minimum, maximum=None, standard=None
    ) -> int | None:
        w = self._wert(zeile, spalten, name)
        if _leer(w):
            if standard is not None:
                return standard
            self._meld("fehler", blatt, nr, titel, f"{titel} fehlt")
            return None
        try:
            z = float(str(w).replace(",", ".")) if isinstance(w, str) else float(w)
        except (TypeError, ValueError):
            z = None
        if z is None or z != int(z) or z < minimum or (maximum is not None and z > maximum):
            grenze = f"{minimum} bis {maximum}" if maximum else f"mindestens {minimum}"
            self._meld(
                "fehler", blatt, nr, titel, f"{titel} muss eine ganze Zahl sein ({grenze}), nicht {w!r}"
            )
            return None
        return int(z)

    def _zeit(self, blatt, nr, zeile, spalten, name, titel) -> time | None:
        w = self._wert(zeile, spalten, name)
        if isinstance(w, datetime):
            return w.time().replace(second=0, microsecond=0)
        if isinstance(w, time):
            return w.replace(second=0, microsecond=0)
        if isinstance(w, (int, float)) and not isinstance(w, bool) and 0 <= w < 1:
            minuten = round(w * 24 * 60)
            return time(minuten // 60 % 24, minuten % 60)
        if isinstance(w, str) and (m := re.fullmatch(r"\s*(\d{1,2})[:.](\d{2})(?::\d{2})?\s*", w)):
            h, mi = int(m[1]), int(m[2])
            if h < 24 and mi < 60:
                return time(h, mi)
        self._meld(
            "fehler",
            blatt,
            nr,
            titel,
            f"{titel} ist keine Uhrzeit (HH:MM): {w!r}" if not _leer(w) else f"{titel} fehlt",
        )
        return None

    # ── Blätter ─────────────────────────────────────────────────────────────
    def _lehrerlisten_blatt(self) -> str | None:
        for name in self.mappe.sheetnames:
            erste = next(iter(self.mappe[name].iter_rows(max_row=1, values_only=True)), None)
            if erste and ist_lehrerliste(list(erste)):
                return name
        return None

    def lese(self) -> Importergebnis:
        blatt = self._lehrerlisten_blatt()
        if blatt is not None:
            zeilen = [list(z) for z in self.mappe[blatt].iter_rows(values_only=True)]
            lese_lehrerliste(blatt, zeilen, self.normalisierung, self.erg)
            if not self._konferenzen():
                self._zeitslots()
            self._verfuegbarkeit()
            return self.erg
        for pflicht in PFLICHTBLAETTER:
            if _norm(pflicht) not in self._blaetter:
                self._meld("fehler", pflicht, 1, "", f"Blatt {pflicht} fehlt")
        if self.erg.fehler:
            return self.erg
        self._themen()
        self._personen()
        self._vorsitze()
        if not self._konferenzen():
            self._zeitslots()
        self._verfuegbarkeit()
        return self.erg

    # ── Konferenztermine (SPEC-0008) ─────────────────────────────────────────
    def _merke(self, a: Aufloesung) -> None:
        for i, bekannt in enumerate(self.erg.aufloesungen):
            if bekannt.schluessel == a.schluessel:
                self.erg.aufloesungen[i] = dataclasses.replace(bekannt, vorkommen=bekannt.vorkommen + 1)
                return
        self.erg.aufloesungen.append(dataclasses.replace(a, vorkommen=1))

    def _konferenz_name(self, blatt, nr, spalte, roh, namen) -> str | None:
        a = self.normalisierung.loese(roh, namen)
        self._merke(a)
        if a.status != SICHER or len(a.konferenzen) != 1:
            hinweis = f" – Vorschlag: {', '.join(a.konferenzen)}" if a.status == VORSCHLAG else ""
            self._meld("fehler", blatt, nr, spalte, f"{roh!r} ist nicht zugeordnet{hinweis}")
            return None
        return a.konferenzen[0]

    def _konferenzen(self) -> bool:
        """Liest das Blatt `Konferenzen`: feste Termine je Konferenz. True, wenn das Blatt vorhanden ist."""
        gelesen = self._zeilen("Konferenzen")
        if gelesen is None:
            return False
        blatt = self._blaetter["konferenzen"]
        if "zeitslots" in self._blaetter:
            self._meld(
                "warnung",
                "Zeitslots",
                1,
                "",
                "Blatt Zeitslots wird ignoriert, die Zeiten stehen im Blatt Konferenzen",
            )
        spalten, daten = gelesen
        namen = [t.name for t in self.erg.themen]
        fenster: dict[str, list[tuple[int, Zeitslot]]] = {}
        werte: dict[str, dict[str, tuple[int, Any]]] = {}
        haupt: dict[str, str] = {}
        reihenfolge: list[Zeitslot] = []

        def merke_wert(name: str, feld: str, nr: int, wert: Any, spalte: str) -> None:
            bisher = werte.setdefault(name, {}).get(feld)
            if bisher is None:
                werte[name][feld] = (nr, wert)
            elif bisher[1] != wert:
                self._meld(
                    "fehler", blatt, nr, spalte, f"{spalte} von {name} weicht von Zeile {bisher[0]} ab"
                )

        for nr, z in daten:
            roh = self._text(blatt, nr, z, spalten, "konferenz", "Konferenz")
            tag = self._text(blatt, nr, z, spalten, "tag", "Tag")
            start = self._zeit(blatt, nr, z, spalten, "start", "Start")
            ende = self._zeit(blatt, nr, z, spalten, "ende", "Ende")
            if roh is None or tag is None or start is None or ende is None:
                continue
            if ende <= start:
                self._meld("fehler", blatt, nr, "Ende", "Ende liegt nicht nach dem Start")
                continue
            name = self._konferenz_name(blatt, nr, "Konferenz", roh, namen)
            if name is None:
                continue
            if name not in namen:
                self._meld(
                    "warnung",
                    blatt,
                    nr,
                    "Konferenz",
                    f"Konferenz {name} hat keine Mitglieder, Termin ignoriert",
                )
                continue
            fenster_neu = Zeitslot(tag, start, ende)
            for andere_nr, w in fenster.get(name, []):
                if w.tag == tag and start < w.ende and w.start < ende:
                    self._meld(
                        "fehler",
                        blatt,
                        nr,
                        "Start",
                        f"Termin überschneidet sich mit Zeile {andere_nr} derselben Konferenz",
                    )
            fenster.setdefault(name, []).append((nr, fenster_neu))
            reihenfolge.append(fenster_neu)
            if "wichtigkeit" in spalten and not _leer(self._wert(z, spalten, "wichtigkeit")):
                w = self._zahl(blatt, nr, z, spalten, "wichtigkeit", "Wichtigkeit", 1, 5)
                if w is not None:
                    merke_wert(name, "wichtigkeit", nr, w, "Wichtigkeit")
            if "mindestanzahl" in spalten and not _leer(self._wert(z, spalten, "mindestanzahl")):
                m = self._zahl(blatt, nr, z, spalten, "mindestanzahl", "Mindestanzahl", 1)
                if m is not None:
                    merke_wert(name, "mindestanzahl", nr, m, "Mindestanzahl")
            oben = self._wert(z, spalten, "hauptkonferenz")
            if not _leer(oben):
                h = self._konferenz_name(blatt, nr, "Hauptkonferenz", str(oben).strip(), namen)
                if h is not None and h not in namen:
                    self._meld(
                        "fehler",
                        blatt,
                        nr,
                        "Hauptkonferenz",
                        f"Hauptkonferenz {h} ist keine bekannte Konferenz",
                    )
                elif h is not None:
                    merke_wert(name, "haupt", nr, h, "Hauptkonferenz")
                    haupt[name] = h
        for name in namen:
            if name not in fenster:
                self._meld(
                    "fehler",
                    blatt,
                    1,
                    "Konferenz",
                    f"Konferenz {name} hat keinen Termin im Blatt Konferenzen",
                )
        self.erg.themen = [
            dataclasses.replace(
                t,
                festtermine=tuple(w for _, w in fenster.get(t.name, [])),
                dauer_min=(fenster[t.name][0][1].dauer_min if t.name in fenster else t.dauer_min),
                **{k: v for k, (_, v) in werte.get(t.name, {}).items()},
            )
            for t in self.erg.themen
        ]
        self.erg.zeitslots = raster_aus_fenstern(reihenfolge)
        self.erg.festtermine = True
        self._pruefe_reihenfolge(fenster, haupt, blatt)
        return True

    def _pruefe_reihenfolge(self, fenster, haupt: dict[str, str], blatt: str) -> None:
        """Warnung, wenn eine Unterkonferenz nicht nach ihrer Hauptkonferenz beginnt (SPEC-0008 FR-08)."""
        tage = list(dict.fromkeys(w.tag for ws in fenster.values() for _, w in ws))
        for unter, ober in haupt.items():
            for nr, u in fenster.get(unter, []):
                if any(
                    (tage.index(u.tag), u.start) < (tage.index(w.tag), w.ende)
                    for _, w in fenster.get(ober, [])
                ):
                    self._meld(
                        "warnung",
                        blatt,
                        nr,
                        "Hauptkonferenz",
                        f"{unter} beginnt nicht nach dem Ende von {ober}",
                    )
                    break

    def _themen(self) -> None:
        gelesen = self._zeilen("Themen")
        if gelesen is None:
            return
        spalten, daten = gelesen
        self.erg.themen_aus_blatt = True
        gesehen: set[str] = set()
        for nr, z in daten:
            name = self._text("Themen", nr, z, spalten, "name", "Name")
            wichtig = self._zahl("Themen", nr, z, spalten, "wichtigkeit", "Wichtigkeit", 1, 5)
            mind = self._zahl("Themen", nr, z, spalten, "mindestanzahl", "Mindestanzahl", 1)
            dauer = self._zahl("Themen", nr, z, spalten, "dauer", "Dauer", 1)
            if name is None or None in (wichtig, mind, dauer):
                continue
            if name in gesehen:
                self._meld("fehler", "Themen", nr, "Name", f"Thema {name} kommt doppelt vor")
                continue
            gesehen.add(name)
            vorsitz = self._wert(z, spalten, "vorsitz")
            vorsitz = None if _leer(vorsitz) else str(vorsitz).strip()
            self.erg.themen.append(Thema(name, wichtig, mind, dauer, vorsitz))
            self._themen_zeilen[name] = nr

    def _vorsitze(self) -> None:
        """Der Vorsitz muss eine Person aus dem Blatt Personen sein (SPEC-0006 FR-06)."""
        namen = {p.name for p in self.erg.personen}
        for t in self.erg.themen:
            if t.vorsitz is not None and t.vorsitz not in namen:
                self._meld(
                    "fehler",
                    "Themen",
                    self._themen_zeilen.get(t.name, 1),
                    "Vorsitz",
                    f"Vorsitz {t.vorsitz} steht nicht im Blatt Personen",
                )

    def _personen(self) -> None:
        spalten, daten = self._zeilen("Personen") or ({}, [])
        bekannt = {t.name for t in self.erg.themen}
        namen: dict[str, int] = {}
        abgeleitet: dict[str, None] = {}
        for nr, z in daten:
            name = self._text("Personen", nr, z, spalten, "name", "Name")
            t1 = self._text("Personen", nr, z, spalten, "thema_1", "Thema 1")
            t2 = self._text("Personen", nr, z, spalten, "thema_2", "Thema 2")
            if name is None or t1 is None or t2 is None:
                continue
            if t1 == t2:
                self._meld("fehler", "Personen", nr, "Thema 2", f"Thema 1 und Thema 2 sind gleich ({t1})")
                continue
            fehlend = [
                (t, s)
                for t, s in ((t1, "Thema 1"), (t2, "Thema 2"))
                if self.erg.themen_aus_blatt and t not in bekannt
            ]
            for t, spalte in fehlend:
                self._meld("fehler", "Personen", nr, spalte, f"Unbekanntes Thema {t}: nicht im Blatt Themen")
            if fehlend:
                continue
            if name in namen:
                self._meld(
                    "warnung",
                    "Personen",
                    nr,
                    "Name",
                    f"Name {name} kommt doppelt vor (erste Zeile {namen[name]} gilt)",
                )
                continue
            namen[name] = nr
            abgeleitet.setdefault(t1)
            abgeleitet.setdefault(t2)
            self.erg.personen.append(Person(name, (t1, t2)))
            self.erg.quelle[("Personen", nr)] = name
        if not self.erg.themen_aus_blatt:
            self.erg.themen = [
                Thema(t, STANDARD_WICHTIGKEIT, STANDARD_MINDESTANZAHL, STANDARD_DAUER_MIN) for t in abgeleitet
            ]

    def _zeitslots(self) -> None:
        gelesen = self._zeilen("Zeitslots")
        if gelesen is None:
            return
        spalten, daten = gelesen
        for nr, z in daten:
            tag = self._text("Zeitslots", nr, z, spalten, "tag", "Tag")
            start = self._zeit("Zeitslots", nr, z, spalten, "start", "Start")
            ende = self._zeit("Zeitslots", nr, z, spalten, "ende", "Ende")
            if tag is None or start is None or ende is None:
                continue
            if ende <= start:
                self._meld("fehler", "Zeitslots", nr, "Ende", "Ende liegt nicht nach dem Start")
                continue
            self.erg.zeitslots.append(Zeitslot(tag, start, ende))

    def _verfuegbarkeit(self) -> None:
        gelesen = self._zeilen("Verfügbarkeiten")
        if gelesen is None:
            return
        spalten, daten = gelesen
        if daten and not self.erg.zeitslots:
            self._meld("fehler", "Verfügbarkeiten", 1, "", "Verfügbarkeiten brauchen das Blatt Zeitslots")
            return
        index = {(s.tag, s.start): i for i, s in enumerate(self.erg.zeitslots)}
        personen = {p.name for p in self.erg.personen}
        for nr, z in daten:
            person = self._text("Verfügbarkeiten", nr, z, spalten, "person", "Person")
            tag = self._text("Verfügbarkeiten", nr, z, spalten, "tag", "Tag")
            start = self._zeit("Verfügbarkeiten", nr, z, spalten, "start", "Start")
            if person is None or tag is None or start is None:
                continue
            if person not in personen:
                self._meld("fehler", "Verfügbarkeiten", nr, "Person", f"Unbekannte Person {person}")
                continue
            if (tag, start) not in index:
                self._meld(
                    "fehler", "Verfügbarkeiten", nr, "Start", f"Kein Zeitslot am {tag} um {start:%H:%M}"
                )
                continue
            self.erg.verfuegbarkeit.setdefault(person, set()).add(index[tag, start])


__all__ = ["BLAETTER", "ImportDateiFehler", "Importergebnis", "Meldung", "ModellFehler", "importiere"]
