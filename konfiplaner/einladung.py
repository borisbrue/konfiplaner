"""Word-Export der Einladung mit Ablaufplan (SPEC-0014)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.exceptions import UnrecognizedImageError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from konfiplaner.excel_export import ExportFehler, ZielDateiExistiertFehler
from konfiplaner.tagesplan import Tagesplan

SPALTEN = ("Zeiten", "Inhalt", "Wer/Wie?")
BREITEN_CM = (4.7, 11.5, 9.7)
KOPFFARBE = "A5A5A5"


def _absatz(container, text: str = "", groesse: float = 12, fett: bool = False, stil: str | None = None):
    p = container.add_paragraph(style=stil) if stil else container.add_paragraph()
    if text:
        lauf = p.add_run(text)
        lauf.bold = fett
        lauf.font.size = Pt(groesse)
    return p


def _fuellung(zelle, farbe: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), farbe)
    zelle._tc.get_or_add_tcPr().append(shd)


def _zelle_fuellen(zelle, zeilen: list[str], titel: str = "") -> None:
    erster = True
    if titel:
        p = zelle.paragraphs[0]
        lauf = p.add_run(titel)
        lauf.bold = True
        lauf.font.size = Pt(12)
        erster = False
    for zeile in zeilen:
        p = zelle.paragraphs[0] if erster else zelle.add_paragraph()
        erster = False
        if zeile:
            p.add_run(zeile).font.size = Pt(12)
    for p in zelle.paragraphs:
        p.paragraph_format.space_after = Pt(2)


def baue_einladung(plan: Tagesplan, konferenzen: list[str], logo: Path | None = None):
    """Erzeugt das Word-Dokument im Speicher (Querformat A4)."""
    dok = Document()
    abschnitt = dok.sections[0]
    abschnitt.orientation = WD_ORIENT.LANDSCAPE
    abschnitt.page_width, abschnitt.page_height = Cm(29.7), Cm(21.0)
    abschnitt.left_margin = abschnitt.right_margin = Cm(1.8)
    abschnitt.top_margin = abschnitt.bottom_margin = Cm(1.4)
    dok.styles["Normal"].font.size = Pt(12)
    dok.styles["Normal"].paragraph_format.space_after = Pt(6)

    if logo is not None and logo.is_file():
        try:
            p = dok.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.add_run().add_picture(str(logo), height=Cm(2.2))
        except (OSError, UnrecognizedImageError):  # unlesbares Bild: ohne Logo weiter
            dok.element.body.remove(p._p)

    kopf = f"{plan.wochentag}, {plan.datum_text}: " if plan.datum else ""
    _absatz(dok, kopf + plan.titel, 26, True)
    if plan.thema:
        _absatz(dok, f"THEMA: {plan.thema}", 18, True)
    if plan.ziele:
        _absatz(dok, "Ziele/Inhalte:", 12, True)
        for ziel in plan.ziele:
            _absatz(dok, ziel, 12, stil="List Bullet")
    for beschriftung, text in (
        ("Arbeitsformen", plan.arbeitsformen),
        ("Vorbereitung", plan.vorbereitung),
        ("Verschriftlichung", plan.verschriftlichung),
    ):
        if text.strip():
            p = _absatz(dok, f"{beschriftung}: ", 12, True)
            p.add_run(text).font.size = Pt(12)

    _absatz(dok, "Ablaufplan" + (f" {plan.datum_text}" if plan.datum else ""), 18, True)
    tabelle = dok.add_table(rows=1, cols=3)
    tabelle.style = "Table Grid"
    tabelle.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabelle.autofit = False
    for zelle, text, breite in zip(tabelle.rows[0].cells, SPALTEN, BREITEN_CM, strict=True):
        zelle.width = Cm(breite)
        _fuellung(zelle, KOPFFARBE)
        lauf = zelle.paragraphs[0].add_run(text)
        lauf.bold = True
        lauf.font.size = Pt(14)
    kopfzeile = tabelle.rows[0]._tr.get_or_add_trPr()
    kopfzeile.append(OxmlElement("w:tblHeader"))
    for block in plan.bloecke:
        zellen = tabelle.add_row().cells
        for zelle, breite in zip(zellen, BREITEN_CM, strict=True):
            zelle.width = Cm(breite)
        _zelle_fuellen(zellen[0], [block.zeit_text])
        inhalt = plan.block_inhalt(block, konferenzen)
        _zelle_fuellen(zellen[1], inhalt.split("\n") if inhalt else [], block.titel)
        _zelle_fuellen(zellen[2], block.wer.split("\n") if block.wer else [])
    return dok


def exportiere_einladung(
    plan: Tagesplan,
    konferenzen: list[str],
    ziel: str | Path,
    ueberschreiben: bool = False,
    logo: Path | None = None,
) -> None:
    """Schreibt die Einladung; Fehler im Plan oder eine ungeprüfte vorhandene Datei brechen ab."""
    fehler = [h.text for h in plan.pruefe() if h.fehler]
    if fehler:
        raise ExportFehler("Der Ablaufplan hat Fehler:\n" + "\n".join(fehler))
    ziel = Path(ziel)
    if ziel.exists() and not ueberschreiben:
        raise ZielDateiExistiertFehler(f"{ziel.name} existiert bereits")
    dok = baue_einladung(plan, konferenzen, logo)
    try:
        dok.save(str(ziel))
    except OSError as e:
        raise ExportFehler(f"{ziel.name} konnte nicht geschrieben werden: {e}") from e
