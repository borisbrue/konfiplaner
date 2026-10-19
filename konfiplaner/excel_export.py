"""Excel-Export des Terminplans (SPEC-0004).

Der Export liest nur den solver-unabhängigen `Plan` (CON-0003) und das `Modell`;
Status, Verstöße, Begründungen und Warnungen werden unverändert übernommen.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from konfiplaner.modell import Modell, Plan, Termin

SPALTEN: dict[str, list[str]] = {
    "Termine": ["Tag", "Start", "Ende", "Thema", "Teilnehmerzahl", "Teilnehmende"],
    "Personen": ["Name", "Tag", "Start", "Ende", "Thema"],
    "Themen": ["Thema", "Wichtigkeit", "Mindestanzahl", "Teilnehmerzahl", "Erfüllt"],
    "Hinweise": ["Art", "Text"],
}
ZEITFORMAT = "HH:MM"


class ExportFehler(Exception):
    """Der Export ist fehlgeschlagen; die Meldung ist für Anwender:innen gedacht."""


class ZielDateiExistiertFehler(ExportFehler):
    """Die Zieldatei existiert und das Überschreiben wurde nicht bestätigt."""


def exportiere(plan: Plan, modell: Modell, ziel: str | Path, ueberschreiben: bool = False) -> None:
    """Schreibt den Plan als `.xlsx`. Vorhandene Dateien werden nur mit `ueberschreiben=True` ersetzt."""
    ziel = Path(ziel)
    if ziel.exists() and not ueberschreiben:
        raise ZielDateiExistiertFehler(f"{ziel.name} existiert bereits")
    sperre = ziel.with_name(f"~${ziel.name}")
    if sperre.exists():
        raise ExportFehler(f"{ziel.name} ist in einem anderen Programm geöffnet")

    mappe = baue_arbeitsmappe(plan, modell)
    # erst in eine Temp-Datei, dann ersetzen: ein Fehler zerstört keine vorhandene Datei
    try:
        fd, temp = tempfile.mkstemp(suffix=".xlsx", dir=ziel.parent)
        os.close(fd)
        try:
            mappe.save(temp)
            os.replace(temp, ziel)
        finally:
            Path(temp).unlink(missing_ok=True)
    except PermissionError as e:
        raise ExportFehler(
            f"{ziel.name} kann nicht geschrieben werden (gesperrt oder keine Berechtigung)"
        ) from e
    except OSError as e:
        raise ExportFehler(f"{ziel.name} kann nicht geschrieben werden: {e.strerror or e}") from e


def tabellen(plan: Plan, modell: Modell) -> dict[str, list[list]]:
    """Datenzeilen je Blatt (ohne Kopfzeile). Grundlage für Excel-Export und GUI-Ansichten."""
    for t in plan.termine:
        if not 0 <= t.start_slot or t.start_slot + t.anzahl_slots > len(modell.zeitslots):
            raise ExportFehler(f"Termin {t.thema} liegt außerhalb der Zeitslots")
    geordnet = sorted(plan.termine, key=lambda t: (t.start_slot, t.thema))
    termine, personen, themen, hinweise = [], [], [], []
    for t in geordnet:
        tag, start, ende = _zeit(modell, t)
        termine.append([tag, start, ende, t.thema, len(t.teilnehmer), ", ".join(t.teilnehmer)])
    for p in modell.personen:
        for t in geordnet:
            if p.name in t.teilnehmer:
                tag, start, ende = _zeit(modell, t)
                personen.append([p.name, tag, start, ende, t.thema])
    for th in modell.themen:
        eigene = [t for t in plan.termine if t.thema == th.name]
        erfuellt = bool(eigene) and all(len(t.teilnehmer) >= th.mindestanzahl for t in eigene)
        themen.append(
            [
                th.name,
                th.wichtigkeit,
                th.mindestanzahl,
                sum(len(t.teilnehmer) for t in eigene),
                "ja" if erfuellt else "nein",
            ]
        )
    if plan.status == "machbar":
        hinweise.append(["Status", "Plan ist nicht optimal"])
    hinweise += [["Begründung", text] for text in plan.begruendung]
    hinweise += [["Verstoß", f"{v.restriktion}: {v.text}"] for v in plan.verstoesse]
    hinweise += [["Warnung", text] for text in plan.warnungen]
    return {"Termine": termine, "Personen": personen, "Themen": themen, "Hinweise": hinweise}


def baue_arbeitsmappe(plan: Plan, modell: Modell) -> Workbook:
    mappe = Workbook()
    mappe.remove(mappe.active)
    for name, zeilen in tabellen(plan, modell).items():
        blatt = mappe.create_sheet(name)
        blatt.append(SPALTEN[name])
        for zelle in blatt[1]:
            zelle.font = Font(bold=True)
        blatt.freeze_panes = "A2"
        for zeile in zeilen:
            blatt.append(zeile)
            for zelle in blatt[blatt.max_row]:
                if hasattr(zelle.value, "hour"):
                    zelle.number_format = ZEITFORMAT
        _breiten(blatt)
    return mappe


def _zeit(modell: Modell, t: Termin):
    erster, letzter = modell.zeitslots[t.start_slot], modell.zeitslots[t.start_slot + t.anzahl_slots - 1]
    return erster.tag, erster.start, letzter.ende


def _breiten(blatt: Worksheet) -> None:
    for spalte in blatt.columns:
        laenge = max(len(str(z.value)) if z.value is not None else 0 for z in spalte)
        blatt.column_dimensions[spalte[0].column_letter].width = min(max(laenge + 2, 10), 60)
